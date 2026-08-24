"""KB 解析器单元测试（Phase B）：覆盖 txt/md/pdf/docx/csv/xlsx + 异常分支。

离线，使用内存字节流构造样本，不依赖外部文件。
"""

import io

import pytest

from src.kb import parsers
from src.kb.parsers import (
    ParseError,
    UnsupportedFormatError,
    parse_bytes,
)


def _build_minimal_pdf(lines) -> bytes:
    """手工拼一个 xref 正确的极简 PDF（ASCII 文本，避免编码问题）。"""
    content = "BT /F1 12 Tf 50 750 Td 14 TL\n"
    for ln in lines:
        content += f"({ln}) Tj T*\n"
    content += "ET"
    content_b = content.encode("latin-1")
    objs = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 595 842]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<<>>",  # 占位，下面替换为 stream
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    objs[3] = (
        b"<<Length " + str(len(content_b)).encode() + b">>\nstream\n"
        + content_b + b"\nendstream"
    )
    pdf = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objs, start=1):
        offsets.append(len(pdf))
        pdf += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(pdf)
    n = len(objs) + 1
    pdf += f"xref\n0 {n}\n".encode()
    pdf += b"0000000000 65535 f \n"
    for off in offsets:
        pdf += f"{off:010d} 00000 n \n".encode()
    pdf += b"trailer\n<</Size " + f"{n}".encode() + b"/Root 1 0 R>>\n"
    pdf += f"startxref\n{xref_pos}\n%%EOF".encode()
    return pdf


def test_supported_extensions_constant():
    assert ".pdf" in parsers.SUPPORTED_EXTENSIONS
    assert ".docx" in parsers.SUPPORTED_EXTENSIONS
    assert ".xlsx" in parsers.SUPPORTED_EXTENSIONS
    assert ".md" in parsers.SUPPORTED_EXTENSIONS


def test_parse_txt_chinese():
    out = parse_bytes("制度.txt", "年假：入职满3年享10天。\n病假：需医院证明。".encode("utf-8"))
    assert "年假" in out and "病假" in out


def test_parse_md_keeps_text():
    md = "# 标题\n\n这是 **粗体** 内容。\n"
    out = parse_bytes("doc.md", md.encode("utf-8"))
    assert "标题" in out and "粗体" in out


def test_parse_csv_as_text():
    csv = "姓名,部门,职位\n张三,研发部,工程师\n".encode("utf-8")
    out = parse_bytes("emp.csv", csv)
    assert "张三" in out and "研发部" in out


def test_parse_pdf():
    pdf = _build_minimal_pdf(["Hello PDF Parser", "Second line"])
    out = parse_bytes("sample.pdf", pdf)
    assert "Hello PDF Parser" in out
    assert "Second line" in out


def test_parse_docx():
    from docx import Document

    doc = Document()
    doc.add_paragraph("员工手册第一章")
    doc.add_paragraph("年假政策：满3年10天")
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "项目"
    tbl.cell(0, 1).text = "标准"
    buf = io.BytesIO()
    doc.save(buf)
    out = parse_bytes("handbook.docx", buf.getvalue())
    assert "员工手册第一章" in out
    assert "年假政策" in out
    assert "项目" in out and "标准" in out


def test_parse_xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "员工"
    ws.append(["工号", "姓名"])
    ws.append(["E001", "张三"])
    buf = io.BytesIO()
    wb.save(buf)
    out = parse_bytes("roster.xlsx", buf.getvalue())
    assert "E001" in out and "张三" in out
    assert "员工" in out  # sheet 标题作为段落


def test_unsupported_format_raises():
    with pytest.raises(UnsupportedFormatError):
        parse_bytes("note.rtf", b"x")
    with pytest.raises(UnsupportedFormatError):
        parse_bytes("noext", b"x")


def test_gbk_text_decoded():
    out = parse_bytes("gbk.txt", "中文GBK内容".encode("gbk"))
    assert "中文GBK内容" in out


def test_parse_pdf_corrupt_raises_parse_error():
    with pytest.raises(ParseError):
        parse_bytes("bad.pdf", b"not a real pdf at all")
