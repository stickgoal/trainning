"""文件解析（Phase B）：把上传的文档抽取为纯文本，供切分与入库。

支持的格式（见 SUPPORTED_EXTENSIONS）：
- 纯文本类：.txt / .md / .markdown / .text / .csv（直接解码，CSV 即文本）
- PDF：pypdf 抽取每页文本
- Word：python-docx 抽取段落
- Excel：openpyxl 抽取每个 sheet 的行列

解析失败或不支持格式时抛明确的异常，由上层（上传接口）转为友好错误。
"""

from io import BytesIO
from pathlib import Path

SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".text", ".csv",
    ".pdf", ".docx", ".xlsx",
}


class UnsupportedFormatError(ValueError):
    """不支持的文件格式。"""


class ParseError(RuntimeError):
    """格式受支持但解析失败（文件损坏等）。"""


def parse_bytes(filename: str, data: bytes) -> str:
    """解析上传文件字节流为纯文本。

    filename 用于推断扩展名；data 为文件原始字节。
    不支持的扩展名抛 UnsupportedFormatError；解析失败抛 ParseError。
    """
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(
            f"不支持的格式：{ext or '(无扩展名)'}。"
            f"支持：{', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    try:
        if ext in (".txt", ".md", ".markdown", ".text", ".csv"):
            return _decode_text(data)
        if ext == ".pdf":
            return _parse_pdf(data)
        if ext == ".docx":
            return _parse_docx(data)
        if ext == ".xlsx":
            return _parse_xlsx(data)
    except UnsupportedFormatError:
        raise
    except Exception as e:  # 解析过程中其它异常统一包装
        raise ParseError(f"解析 {filename} 失败：{e}") from e

    raise UnsupportedFormatError(f"未处理的扩展名：{ext}")


def _decode_text(data: bytes) -> str:
    """文本类：优先 utf-8-sig，回退 gbk，最后 utf-8。"""
    for enc in ("utf-8-sig", "gbk", "utf-8"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def _parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    parts = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text:
            parts.append(text)
    if not parts:
        return ""
    return "\n\n".join(parts)


def _parse_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    # 也抽取表格内容
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts)


def _parse_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    blocks: list[str] = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in r]
            line = "\t".join(cells).strip()
            if line:
                rows.append(line)
        if rows:
            blocks.append(f"# {ws.title}\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(blocks)
